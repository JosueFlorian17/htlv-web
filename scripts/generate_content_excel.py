import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_content_collection_workbook(output_path):
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles & Colors (Consortium Burgundy / Slate Medical Palette)
    PRIMARY_COLOR = "5B0617"  # #5b0617 Burgundy
    PRIMARY_LIGHT = "FFDADA"  # Light Burgundy
    HEADER_FILL = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    SUBHEADER_FILL = PatternFill(start_color="1D2B3A", end_color="1D2B3A", fill_type="solid") # Dark Navy
    SUBHEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    TITLE_FONT = Font(name="Calibri", size=14, bold=True, color=PRIMARY_COLOR)
    SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="564242")

    FILL_PLACEHOLDER = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
    FILL_INPUT = PatternFill(start_color="FFFDE7", end_color="FFFDE7", fill_type="solid") # Soft yellow to highlight input area
    FILL_GUIDE = PatternFill(start_color="E8F4F8", end_color="E8F4F8", fill_type="solid") # Soft blue/gray for guidelines

    BORDER_THIN = Border(
        left=Side(style='thin', color='DCC0C0'),
        right=Side(style='thin', color='DCC0C0'),
        top=Side(style='thin', color='DCC0C0'),
        bottom=Side(style='thin', color='DCC0C0')
    )
    BORDER_BOTTOM_DOUBLE = Border(
        bottom=Side(style='double', color='5B0617')
    )

    ALIGN_LEFT = Alignment(horizontal='left', vertical='top', wrap_text=True)
    ALIGN_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ALIGN_HEADER = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # -------------------------------------------------------------
    # HOJA 0: INSTRUCCIONES
    # -------------------------------------------------------------
    ws_intro = wb.create_sheet(title="00_INSTRUCCIONES_Y_GUIA")
    ws_intro.views.sheetView[0].showGridLines = True

    intro_headers = [
        "GUÍA DE RECOLECCIÓN Y VALIDACIÓN DE CONTENIDOS OFICIALES",
        "RED INTERNACIONAL DE INVESTIGACIÓN E INNOVACIÓN EN HTLV (RIII-HTLV)",
        "Este documento servirá para sustituir toda la información de maqueta (placeholders) de la plataforma web por el contenido científico e institucional definitivo."
    ]

    ws_intro.cell(row=1, column=1, value=intro_headers[0]).font = Font(name="Calibri", size=15, bold=True, color=PRIMARY_COLOR)
    ws_intro.cell(row=2, column=1, value=intro_headers[1]).font = Font(name="Calibri", size=12, bold=True, color="1D2B3A")
    ws_intro.cell(row=3, column=1, value=intro_headers[2]).font = Font(name="Calibri", size=10, italic=True, color="564242")

    instructions_data = [
        ["PASO", "INSTRUCCIÓN / CÓDIGO DE COLORES", "DESCRIPCIÓN Y BUENAS PRÁCTICAS"],
        ["1", "Revisar cada Pestaña Temática", "El libro está dividido en hojas: 01_INICIO, 02_ACERCA_DE_HTLV, 03_RED_CONSORCIO, 04_INVESTIGACION, 05_RECURSOS_MULTIMEDIA, 06_OPORTUNIDADES, 07_CONTACTO_Y_LEGAL."],
        ["2", "Columna C: Texto Actual (Placeholder)", "Muestra exactamente la redacción provisional que se encuentra actualmente en la web para su referencia."],
        ["3", "Columna D: Texto Oficial Definitivo (Celdas Amarillas)", "✍️ AQUÍ DEBE ESCRIBIR SU TEXTO DEFINITIVO. Las celdas en amarillo suave están listas para redactar o pegar la información oficial validada."],
        ["4", "Columna E: Pautas y Recomendaciones", "Lea las sugerencias de redacción para mantener el rigor biomédico, el tono institucional y los puntos clave a destacar."],
        ["5", "Columna F: Límite de Caracteres / Palabras", "Respete las longitudes sugeridas para asegurar que el diseño visual, las tarjetas y los menús mantengan una diagramación estética óptima tanto en PC como en celulares."],
        ["6", "Columna G: Estado de Revisión", "Marque el estado de cada fila como: 'Pendiente', 'En Revisión' o 'Aprobado'."],
        ["7", "Enlaces, Imágenes y PDFs", "Si el contenido requiere un enlace externo, archivo descargable o foto institucional, por favor pegue la URL o el nombre exacto del archivo adjunto."],
        ["8", "Respaldo Legal & Ético", "Toda la información médica debe estar alineada a la Ley N° 29733 (Protección de Datos Personales) y recomendaciones internacionales de consenso."]
    ]

    for r_idx, row in enumerate(instructions_data, start=5):
        for c_idx, val in enumerate(row, start=1):
            cell = ws_intro.cell(row=r_idx, column=c_idx, value=val)
            cell.border = BORDER_THIN
            if r_idx == 5:
                cell.fill = SUBHEADER_FILL
                cell.font = SUBHEADER_FONT
                cell.alignment = ALIGN_HEADER
            else:
                cell.font = Font(name="Calibri", size=10, bold=(c_idx == 1 or c_idx == 2))
                cell.alignment = ALIGN_LEFT
                if c_idx == 1:
                    cell.alignment = ALIGN_CENTER
                if r_idx % 2 == 0:
                    cell.fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    ws_intro.column_dimensions['A'].width = 10
    ws_intro.column_dimensions['B'].width = 38
    ws_intro.column_dimensions['C'].width = 85

    # Helper function to create content sheet
    def create_content_sheet(sheet_title, page_title, page_subtitle, rows_data):
        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True

        # Header Title
        ws.cell(row=1, column=1, value=page_title).font = TITLE_FONT
        ws.cell(row=2, column=1, value=page_subtitle).font = SUBTITLE_FONT

        headers = [
            "ID_TECNICO",
            "SECCIÓN / COMPONENTE",
            "TEXTO ACTUAL EN LA WEB (PLACEHOLDER)",
            "TEXTO OFICIAL / DEFINITIVO (ESCRIBIR AQUÍ)",
            "PAUTAS Y RECOMENDACIONES DE CONTENIDO",
            "LONGITUD SUGERIDA",
            "ESTADO"
        ]

        # Table Header (Row 4)
        for c_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=c_idx, value=h)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = ALIGN_HEADER
            cell.border = BORDER_THIN

        # Data Rows
        for r_idx, row in enumerate(rows_data, start=5):
            for c_idx, val in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.border = BORDER_THIN
                cell.alignment = ALIGN_LEFT

                if c_idx == 1:  # ID
                    cell.font = Font(name="Calibri", size=9.5, bold=True, color="555F6F")
                    cell.alignment = ALIGN_CENTER
                    cell.fill = FILL_PLACEHOLDER
                elif c_idx == 2:  # Seccion
                    cell.font = Font(name="Calibri", size=10, bold=True, color="191C1E")
                elif c_idx == 3:  # Placeholder
                    cell.font = Font(name="Calibri", size=9.5, color="564242")
                    cell.fill = FILL_PLACEHOLDER
                elif c_idx == 4:  # Campo en Blanco
                    cell.font = Font(name="Calibri", size=10, color="000000")
                    cell.fill = FILL_INPUT
                elif c_idx == 5:  # Pautas
                    cell.font = Font(name="Calibri", size=9.5, italic=True, color="2C3E50")
                    cell.fill = FILL_GUIDE
                elif c_idx == 6:  # Longitud
                    cell.font = Font(name="Calibri", size=9.5, bold=True, color="5B0617")
                    cell.alignment = ALIGN_CENTER
                elif c_idx == 7:  # Estado
                    cell.font = Font(name="Calibri", size=9.5, bold=True, color="897172")
                    cell.alignment = ALIGN_CENTER

            ws.row_dimensions[r_idx].height = 45

        # Column widths
        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 28
        ws.column_dimensions['C'].width = 45
        ws.column_dimensions['D'].width = 50
        ws.column_dimensions['E'].width = 45
        ws.column_dimensions['F'].width = 18
        ws.column_dimensions['G'].width = 14
        ws.freeze_panes = "A5"
        return ws

    # -------------------------------------------------------------
    # HOJA 1: INICIO (HOME)
    # -------------------------------------------------------------
    home_rows = [
        [
            "HOME_HERO_TAG",
            "Hero - Insignia Superior",
            "Liderazgo Científico Global • UPCH • IMTAvH • Perú",
            "",
            "Identificar las entidades coordinadoras oficiales y el alcance internacional.",
            "Máx. 50 caracteres",
            "Pendiente"
        ],
        [
            "HOME_HERO_TITLE",
            "Hero - Título Principal",
            "Red Internacional de Investigación e Innovación en HTLV",
            "",
            "Nombre oficial del consorcio/red o lema institucional de alto impacto.",
            "Máx. 70 caracteres",
            "Pendiente"
        ],
        [
            "HOME_HERO_SUBTITLE",
            "Hero - Subtítulo / Bajada",
            "Conectando a centros de excelencia en retrovirología humana para acelerar diagnósticos tempranos, protocolos terapéuticos y la erradicación global de las patologías asociadas al HTLV-1 y HTLV-2.",
            "",
            "Declaración concisa del propósito de la red, impacto esperado y retrovirología.",
            "150 - 220 caracteres",
            "Pendiente"
        ],
        [
            "HOME_HERO_CTA_1",
            "Hero - Botón de Acción 1",
            "¿Qué es el HTLV? (Guía Clínica)",
            "",
            "Texto del botón principal hacia la guía informativa.",
            "Máx. 35 caracteres",
            "Pendiente"
        ],
        [
            "HOME_HERO_CTA_2",
            "Hero - Botón de Acción 2",
            "Líneas de Investigación",
            "",
            "Texto del botón hacia proyectos científicos.",
            "Máx. 30 caracteres",
            "Pendiente"
        ],
        [
            "HOME_HERO_CTA_3",
            "Hero - Botón de Acción 3",
            "Repositorio Científico",
            "",
            "Texto del botón hacia la biblioteca de publicaciones.",
            "Máx. 30 caracteres",
            "Pendiente"
        ],
        [
            "HOME_AUTH_TAG",
            "Autoridad - Etiqueta",
            "Consorcio Institucional",
            "",
            "Categoría de la sección de autoridad.",
            "Máx. 30 caracteres",
            "Pendiente"
        ],
        [
            "HOME_AUTH_TITLE",
            "Autoridad - Título de Sección",
            "Excelencia Científica desde el Foco Endémico hacia el Mundo",
            "",
            "Título enfocado en la relevancia del Perú como centro endémico de referencia mundial.",
            "Máx. 65 caracteres",
            "Pendiente"
        ],
        [
            "HOME_AUTH_P1",
            "Autoridad - Párrafo 1",
            "La RIII-HTLV es una red colaborativa internacional creada para unir la experiencia clínica de zonas de alta prevalencia como el Perú con la capacidad de investigación básica y molecular de laboratorios de referencia en Europa, América Latina y Asia.",
            "",
            "Explicar el origen de la alianza, nodos internacionales y alcance colaborativo.",
            "200 - 300 caracteres",
            "Pendiente"
        ],
        [
            "HOME_AUTH_P2",
            "Autoridad - Párrafo 2",
            "Coordinada con la participación del Instituto de Medicina Tropical Alexander von Humboldt de la Universidad Peruana Cayetano Heredia (UPCH), nuestra plataforma promueve la estandarización de pruebas diagnósticas, cohortes clínicas longitudinales y políticas de salud pública preventivas.",
            "",
            "Mencionar la coordinación del IMTAvH/UPCH, estándares diagnósticos y cohortes.",
            "200 - 300 caracteres",
            "Pendiente"
        ],
        [
            "HOME_AUTH_QUOTE",
            "Autoridad - Cita Destacada",
            "\"La integración de datos clínicos y genómicos salva vidas.\" — Comité Científico RIII-HTLV",
            "",
            "Frase inspiradora o principio rector atribuido al Comité Científico o Presidente de la Red.",
            "Máx. 120 caracteres",
            "Pendiente"
        ],
        [
            "HOME_BENTO_01",
            "Pilar 1 - Sobre HTLV",
            "Título: Sobre HTLV | Descripción: Guía médica y clínica",
            "",
            "Módulo de acceso a la guía de patogénesis y espectro clínico.",
            "Título: 2-3 palabras\nDesc: 4-6 palabras",
            "Pendiente"
        ],
        [
            "HOME_BENTO_02",
            "Pilar 2 - La Red",
            "Título: La Red | Descripción: Instituciones y miembros",
            "",
            "Módulo hacia las instituciones y mapa de miembros.",
            "Título: 2-3 palabras\nDesc: 4-6 palabras",
            "Pendiente"
        ],
        [
            "HOME_BENTO_03",
            "Pilar 3 - Investigación",
            "Título: Investigación | Descripción: Proyectos y cohortes",
            "",
            "Módulo hacia las líneas y ensayos clínicos.",
            "Título: 2-3 palabras\nDesc: 4-6 palabras",
            "Pendiente"
        ],
        [
            "HOME_BENTO_04",
            "Pilar 4 - Repositorio",
            "Título: Repositorio | Descripción: Papers y publicaciones",
            "",
            "Módulo hacia artículos indexados en PubMed/Scopus.",
            "Título: 2-3 palabras\nDesc: 4-6 palabras",
            "Pendiente"
        ],
        [
            "HOME_BENTO_05",
            "Pilar 5 - Recursos",
            "Título: Recursos | Descripción: Multimedia y guías PDF",
            "",
            "Módulo hacia videos interactivos, modelos 3D y calculadoras.",
            "Título: 2-3 palabras\nDesc: 4-6 palabras",
            "Pendiente"
        ],
        [
            "HOME_BENTO_06",
            "Pilar 6 - Foro",
            "Título: Foro Científico | Descripción: Discusión académica",
            "",
            "Módulo de debate entre profesionales de la salud e investigadores.",
            "Título: 2-3 palabras\nDesc: 4-6 palabras",
            "Pendiente"
        ],
        [
            "HOME_STATS_1",
            "Métrica 1 - Países",
            "Cifra: 12+ | Etiqueta: Países y Centros Conectados",
            "",
            "Número real o estimado de países representados en la red.",
            "Cifra breve + etiqueta",
            "Pendiente"
        ],
        [
            "HOME_STATS_2",
            "Métrica 2 - Cohortes",
            "Cifra: 35+ | Etiqueta: Años de Investigación Continua",
            "",
            "Años de trayectoria de investigación clínica en UPCH/IMTAvH o cohortes activas.",
            "Cifra breve + etiqueta",
            "Pendiente"
        ],
        [
            "HOME_STATS_3",
            "Métrica 3 - Publicaciones",
            "Cifra: 150+ | Etiqueta: Publicaciones Indexadas",
            "",
            "Total aproximado de artículos científicos publicados por los miembros.",
            "Cifra breve + etiqueta",
            "Pendiente"
        ],
        [
            "HOME_STATS_4",
            "Métrica 4 - Pacientes/Muestras",
            "Cifra: 5,000+ | Etiqueta: Pacientes y Muestras en Seguimiento",
            "",
            "Registro de base de datos o biobanco disponible.",
            "Cifra breve + etiqueta",
            "Pendiente"
        ]
    ]
    create_content_sheet("01_INICIO", "PÁGINA PRINCIPAL (HOME)", "Textos de la página de inicio, secciones de impacto, pilares y métricas estadísticas.", home_rows)

    # -------------------------------------------------------------
    # HOJA 2: ACERCA DE HTLV (ABOUT)
    # -------------------------------------------------------------
    about_rows = [
        [
            "ABOUT_HEADER_TITLE",
            "Cabecera - Título Principal",
            "Virus Linfotrópico de Células T Humanas (HTLV)",
            "",
            "Título de la guía médica oficial sobre HTLV-1 y HTLV-2.",
            "Máx. 60 caracteres",
            "Pendiente"
        ],
        [
            "ABOUT_HEADER_SUBTITLE",
            "Cabecera - Bajada de Autoridad",
            "Consenso científico, epidemiología global y protocolos clínicos elaborados por investigadores de la Universidad Peruana Cayetano Heredia (UPCH) y la Red Internacional RIII-HTLV.",
            "",
            "Texto de respaldo institucional, comité editorial y fecha de actualización.",
            "150 - 250 caracteres",
            "Pendiente"
        ],
        [
            "ABOUT_SEC1_DEF_P1",
            "1. ¿Qué es el HTLV? - Definición Virológica",
            "El Virus Linfotrópico de Células T Humanas (HTLV-1 y HTLV-2) es un retrovirus complejo de ARN perteneciente a la familia Retroviridae y al género Deltaretrovirus. Descubierto en 1980 por el equipo del Dr. Robert Gallo, fue el primer retrovirus oncogénico humano identificado.",
            "",
            "Definición taxonómica, clasificación de retrovirus, descubrimiento e historia científica.",
            "200 - 350 caracteres",
            "Pendiente"
        ],
        [
            "ABOUT_SEC1_DEF_P2",
            "1. ¿Qué es el HTLV? - Mecanismo de Integración",
            "A diferencia de otros virus, el HTLV se integra en el genoma de los linfocitos T CD4+ y CD8+ en forma de ADN proviral, estableciendo una infección de por vida. La replicación ocurre principalmente por expansión clonal de las células infectadas mediada por las proteínas reguladoras Tax y HBZ.",
            "",
            "Explicación de latencia proviral, expansión clonal y genes clave (Tax, HBZ, Rex).",
            "200 - 350 caracteres",
            "Pendiente"
        ],
        [
            "ABOUT_SEC2_SPEC_ASYMPT",
            "2. Espectro Clínico - Portadores Asintomáticos (~90%)",
            "Aproximadamente el 90% de las personas seropositivas no desarrollan síntomas clínicos a lo largo de su vida, pero pueden transmitir el virus de forma inadvertida. Requieren consejería especializada y control de carga proviral.",
            "",
            "Datos precisos sobre el estado de portador, seguimiento serológico y manejo de la ansiedad.",
            "150 - 250 caracteres",
            "Pendiente"
        ],
        [
            "ABOUT_SEC2_SPEC_HAM_TSP",
            "2. Espectro Clínico - Mielopatía / TSP (~2-5%)",
            "Mielopatía Asociada al HTLV-1 / Paraparesia Espástica Tropical (HAM/TSP): Trastorno neurológico crónico progresivo caracterizado por debilidad en extremidades inferiores, espasticidad, hiperreflexia y disfunción esfinteriana.",
            "",
            "Cuadro clínico de HAM/TSP, criterios de Osame, diagnóstico diferencial y manejo de soporte.",
            "180 - 300 caracteres",
            "Pendiente"
        ],
        [
            "ABOUT_SEC2_SPEC_ATL",
            "2. Espectro Clínico - Leucemia ATL (~3-5%)",
            "Leucemia/Linfoma de Células T del Adulto (ATL): Neoplasia hematológica agresiva de linfocitos T maduros CD4+CD25+. Presenta subtipos: agudo, linfomatoso, crónico y smoldering según los criterios de Shimoyama.",
            "",
            "Subtipos clínicos de Shimoyama, pronóstico, terapias antivirales e inmunoterapia (Mogamulizumab).",
            "180 - 300 caracteres",
            "Pendiente"
        ],
        [
            "ABOUT_SEC2_SPEC_OTHER",
            "2. Espectro Clínico - Manifestaciones Inflamatorias",
            "Uveítis asociada a HTLV-1 (HAU), dermatitis infectiva pediátrica, síndrome de Sjögren, artritis y mayor susceptibilidad a coinfecciones como Strongyloides stercoralis y Tuberculosis pulmonar.",
            "",
            "Detallar la asociación con estrongiloidiasis severa, sarna costrosa y tuberculosis.",
            "180 - 300 caracteres",
            "Pendiente"
        ],
        [
            "ABOUT_SEC3_TRANS_VERT",
            "3. Transmisión - Vía Vertical / Lactancia",
            "Transmisión Materno-Infantil: Principalmente a través de la lactancia materna prolongada (> 6 meses). Se previene mediante tamizaje serológico prenatal y reemplazo con sucedáneos de leche materna (fórmula).",
            "",
            "Protocolo obstétrico y pediátrico para prevención de transmisión vertical.",
            "150 - 250 caracteres",
            "Pendiente"
        ],
        [
            "ABOUT_SEC3_TRANS_SEX",
            "3. Transmisión - Vía Sexual",
            "Relaciones sexuales sin protección: Mayor eficiencia de transmisión de hombre a mujer. El uso correcto y consistente del preservativo previene el contagio.",
            "",
            "Pautas de prevención en parejas serodiscordantes y educación sexual.",
            "150 - 250 caracteres",
            "Pendiente"
        ],
        [
            "ABOUT_SEC3_TRANS_PAREN",
            "3. Transmisión - Vía Parenteral / Sangre",
            "Transfusión sanguínea y trasplante de órganos sin tamizaje previo, o uso compartido de jeringas. Se controla con el tamizaje serológico obligatorio en el 100% de los bancos de sangre.",
            "",
            "Marco regulatorio de tamizaje en bancos de sangre y bioseguridad.",
            "150 - 250 caracteres",
            "Pendiente"
        ],
        [
            "ABOUT_SEC4_PERU_ROLE",
            "4. Liderazgo de Perú y UPCH",
            "El Perú es una de las áreas endémicas más estudiadas del mundo. El Instituto de Medicina Tropical Alexander von Humboldt (IMTAvH) de la UPCH ha liderado por más de 3 décadas estudios pioneros en epidemiología, diagnóstico molecular y coinfecciones.",
            "",
            "Resaltar las contribuciones científicas históricas de la UPCH/IMTAvH a la retrovirología mundial.",
            "200 - 350 caracteres",
            "Pendiente"
        ],
        [
            "ABOUT_SEC5_DIAG_ELISA",
            "5. Diagnóstico - Tamizaje Serológico",
            "Pruebas de Tamizaje (ELISA / Quimioluminiscencia): Detectan anticuerpos contra antígenos de la envoltura (gp46, gp21) y del core viral (p19, p24). Tienen alta sensibilidad.",
            "",
            "Especificaciones de pruebas serológicas de 4ta generación y algoritmos de tamizaje.",
            "150 - 250 caracteres",
            "Pendiente"
        ],
        [
            "ABOUT_SEC5_DIAG_CONF",
            "5. Diagnóstico - Pruebas Confirmatorias",
            "Pruebas Confirmatorias y Diferenciación (Western Blot, Line Immunoassay INNO-LIA y PCR en tiempo real): Permiten confirmar la infección y diferenciar con exactitud entre HTLV-1 y HTLV-2.",
            "",
            "Criterios de positividad de Western Blot (OMS/CDC) y diferenciación HTLV-1 vs HTLV-2.",
            "150 - 250 caracteres",
            "Pendiente"
        ],
        [
            "ABOUT_SEC5_DIAG_PROVIRAL",
            "5. Diagnóstico - Carga Proviral (qPCR)",
            "Cuantificación de Carga Proviral (qPCR): Determina el porcentaje de células mononucleares de sangre periférica infectadas (copias provirales / 100 PBMC). Cargas > 1-4% se asocian a mayor riesgo de HAM/TSP.",
            "",
            "Punto de corte de carga proviral, valor pronóstico y utilidad en seguimiento clínico.",
            "150 - 250 caracteres",
            "Pendiente"
        ],
        [
            "ABOUT_FAQ_1",
            "6. FAQ 1 - ¿HTLV es igual al VIH?",
            "Pregunta: ¿El HTLV-1 es lo mismo que el VIH?\nRespuesta: No. Aunque ambos son retrovirus, el HTLV no produce SIDA. En lugar de destruir células T, induce proliferación celular y en raros casos cuadros inflamatorios (HAM/TSP) o neoplásicos (ATL).",
            "",
            "Aclaración para pacientes y público general sobre las diferencias fundamentales con el VIH.",
            "Pregunta: 1 línea\nRespuesta: 200-300 caracteres",
            "Pendiente"
        ],
        [
            "ABOUT_FAQ_2",
            "6. FAQ 2 - ¿Existe cura o vacuna?",
            "Pregunta: ¿Existe cura o vacuna para el HTLV?\nRespuesta: Actualmente no existe vacuna disponible ni fármaco que erradique el virus del genoma. El manejo médico se enfoca en prevención, monitoreo periódico y tratamiento sintomático temprano.",
            "",
            "Explicación honesta sobre el estado actual de la investigación de vacunas y antivirales.",
            "Pregunta: 1 línea\nRespuesta: 200-300 caracteres",
            "Pendiente"
        ],
        [
            "ABOUT_FAQ_3",
            "6. FAQ 3 - Prevención Madre-Hijo",
            "Pregunta: ¿Cómo se previene la transmisión de madre a hijo?\nRespuesta: La principal medida es el despistaje prenatal en el embarazo. En madres seropositivas, se recomienda suspender la lactancia materna y alimentar al bebé con fórmula maternizada.",
            "",
            "Guía para gestantes y pediatras.",
            "Pregunta: 1 línea\nRespuesta: 200-300 caracteres",
            "Pendiente"
        ],
        [
            "ABOUT_FAQ_4",
            "6. FAQ 4 - Dónde hacerse la prueba en Perú",
            "Pregunta: ¿Dónde puedo realizarme una prueba de descarte en el Perú?\nRespuesta: En bancos de sangre autorizados, hospitales del MINSA, EsSalud y en el Instituto de Medicina Tropical Alexander von Humboldt (UPCH) en Lima.",
            "",
            "Directorio referencial de laboratorios de referencia en Perú y regiones.",
            "Pregunta: 1 línea\nRespuesta: 200-300 caracteres",
            "Pendiente"
        ]
    ]
    create_content_sheet("02_ACERCA_DE_HTLV", "GUÍA CLÍNICA: ACERCA DEL HTLV", "Definiciones biomédicas, espectro clínico (90/10), vías de transmisión, diagnóstico y preguntas frecuentes.", about_rows)

    # -------------------------------------------------------------
    # HOJA 3: RED Y CONSORCIO (NETWORK)
    # -------------------------------------------------------------
    net_rows = [
        [
            "NET_HEADER_TITLE",
            "Cabecera - Título de la Red",
            "La Red Internacional de Investigación e Innovación en HTLV (RIII-HTLV)",
            "",
            "Nombre institucional oficial de la red y estructura de gobernanza.",
            "Máx. 80 caracteres",
            "Pendiente"
        ],
        [
            "NET_HEADER_DESC",
            "Cabecera - Descripción General",
            "Una alianza científica global que conecta a centros de investigación líderes de América Latina, Europa y Asia para estandarizar diagnósticos, acelerar descubrimientos terapéuticos y promover la salud pública.",
            "",
            "Misión, visión institucional y alcance geopolítico del consorcio.",
            "180 - 280 caracteres",
            "Pendiente"
        ],
        # Instituciones
        [
            "NET_INST_01",
            "Institución 1 - Sede Principal",
            "Nombre: Universidad Peruana Cayetano Heredia (UPCH)\nUnidad: Instituto de Medicina Tropical Alexander von Humboldt (IMTAvH)\nPaís: Perú\nRol: Centro de Coordinación & Liderazgo Clínico-Epidemiológico\nEnfoque: Cohortes de seguimiento en HAM/TSP, coinfección TB-HTLV y diagnóstico molecular.\nWeb: https://imtavh.cayetano.edu.pe",
            "",
            "Datos oficiales de la sede de coordinación central en Lima, Perú.",
            "Completar todos los campos del formato institucional",
            "Pendiente"
        ],
        [
            "NET_INST_02",
            "Institución 2 - Nodo Europeo (UK)",
            "Nombre: Imperial College London\nUnidad: Section of Virology & National Centre for Human Retrovirology\nPaís: Reino Unido\nRol: Centro de Biología Molecular e Inmunología\nEnfoque: Mecanismos de persistencia viral, carga proviral y regulación inmunológica de células T.\nWeb: https://www.imperial.ac.uk",
            "",
            "Centro colaborador en Reino Unido.",
            "Completar todos los campos del formato institucional",
            "Pendiente"
        ],
        [
            "NET_INST_03",
            "Institución 3 - Nodo Internacional (Francia)",
            "Nombre: Institut Pasteur\nUnidad: Unité d’Epidémiologie et Physiopathologie des Virus Oncogènes\nPaís: Francia\nRol: Investigación Oncológica y Patogénesis de ATL\nEnfoque: Oncogénesis viral del gen Tax/HBZ y transformación celular en linfocitos T.\nWeb: https://www.pasteur.fr",
            "",
            "Centro colaborador en Francia.",
            "Completar todos los campos del formato institucional",
            "Pendiente"
        ],
        [
            "NET_INST_04",
            "Institución 4 - Nodo Sudamericano (Brasil)",
            "Nombre: Universidade de São Paulo (USP) / Fiocruz\nUnidad: Instituto Gonçalo Moniz & FMUSP\nPaís: Brasil\nRol: Vigilancia Epidemiológica y Ensayos Multicéntricos\nEnfoque: Estudios de prevalencia comunitaria, impacto en donantes de sangre y salud pública.\nWeb: https://www.fiocruz.br",
            "",
            "Centro colaborador en Brasil.",
            "Completar todos los campos del formato institucional",
            "Pendiente"
        ],
        [
            "NET_INST_05",
            "Institución 5 - Nodo Ibérico (España)",
            "Nombre: Instituto de Salud Carlos III (ISCIII)\nUnidad: Centro Nacional de Microbiología - Madrid\nPaís: España\nRol: Registro Nacional y Redes de Vigilancia\nEnfoque: Monitoreo de subtipos virales y transmisión transfronteriza en Europa y América Latina.\nWeb: https://www.isciii.es",
            "",
            "Centro colaborador en España.",
            "Completar todos los campos del formato institucional",
            "Pendiente"
        ],
        [
            "NET_INST_06",
            "Institución 6 - Nodo Asia-Pacífico (Japón)",
            "Nombre: St. Marianna University School of Medicine\nUnidad: Department of Rare Diseases & Neurovirology\nPaís: Japón\nRol: Centro de Excelencia en Terapias para HAM/TSP\nEnfoque: Ensayos clínicos avanzados de anticuerpos monoclonales y biomarcadores neurológicos.\nWeb: https://www.marianna-u.ac.jp",
            "",
            "Centro colaborador en Japón.",
            "Completar todos los campos del formato institucional",
            "Pendiente"
        ],
        [
            "NET_INST_07_NUEVA",
            "Institución 7 - (Añadir si aplica)",
            "[Espacio para nueva institución miembro o laboratorio afiliado]",
            "",
            "Indicar si hay más universidades, hospitales o institutos que deban agregarse.",
            "Completar formato institucional",
            "Pendiente"
        ],
        # Comité Científico
        [
            "NET_DIR_01",
            "Directivo 1 - Presidente / Líder",
            "Nombre: Dr. Eduardo Gotuzzo\nCargo: Presidente Honorario | IMTAvH - UPCH (Perú)\nBiografía: Profesor emérito e investigador pionero en enfermedades infecciosas y tropicales en América Latina.\nFoto: [Enlace o archivo adjunto]",
            "",
            "Nombre completo, afiliación académica oficial, síntesis curricular de 2 líneas y foto.",
            "Biografía: 120 - 180 caracteres",
            "Pendiente"
        ],
        [
            "NET_DIR_02",
            "Directivo 2 - Vicepresidente / Co-líder",
            "Nombre: Prof. Graham P. Taylor\nCargo: Vicepresidente | Imperial College London (UK)\nBiografía: Líder del Centro Nacional de Retrovirología Humana de Reino Unido y referente en manejo de ATL y HAM/TSP.\nFoto: [Enlace o archivo adjunto]",
            "",
            "Nombre completo, afiliación académica oficial, síntesis curricular de 2 líneas y foto.",
            "Biografía: 120 - 180 caracteres",
            "Pendiente"
        ],
        [
            "NET_DIR_03",
            "Directivo 3 - Coordinadora Científica",
            "Nombre: Dra. Elsa González\nCargo: Coordinadora Científica | IMTAvH - UPCH (Perú)\nBiografía: Investigadora principal en diagnóstico molecular, coinfecciones y vigilancia materno-infantil de HTLV-1.\nFoto: [Enlace o archivo adjunto]",
            "",
            "Nombre completo, afiliación académica oficial, síntesis curricular de 2 líneas y foto.",
            "Biografía: 120 - 180 caracteres",
            "Pendiente"
        ],
        [
            "NET_DIR_04_NUEVO",
            "Directivo 4 - (Añadir Investigador)",
            "[Espacio para nuevo miembro del comité científico o directivo]",
            "",
            "Indicar miembros adicionales del comité de bioética, publicaciones o investigación.",
            "Biografía: 120 - 180 caracteres",
            "Pendiente"
        ],
        [
            "NET_JOIN_INFO",
            "Adhesión a la Red - Cómo Unirse",
            "Texto explicativo sobre los requisitos para que nuevos laboratorios e investigadores soliciten su incorporación a la red RIII-HTLV.",
            "",
            "Detallar el proceso de postulación, criterios de elegibilidad y correo de contacto.",
            "180 - 280 caracteres",
            "Pendiente"
        ]
    ]
    create_content_sheet("03_RED_CONSORCIO", "RED INTERNACIONAL & MIEMBROS", "Fichas institucionales de miembros, comité científico/directivo y directrices de incorporación a la red.", net_rows)

    # -------------------------------------------------------------
    # HOJA 4: INVESTIGACIÓN (RESEARCH)
    # -------------------------------------------------------------
    res_rows = [
        [
            "RES_HEADER_TITLE",
            "Cabecera - Título de Investigación",
            "Líneas de Investigación, Ensayos Clínicos y Proyectos de la Red",
            "",
            "Título principal de la sección de investigación científica.",
            "Máx. 75 caracteres",
            "Pendiente"
        ],
        [
            "RES_HEADER_DESC",
            "Cabecera - Descripción de Enfoque",
            "Desarrollando investigación traslacional de frontera desde la biología molecular hasta la formulación de políticas públicas para la erradicación del HTLV.",
            "",
            "Resumen del enfoque traslacional (básica, clínica, epidemiológica y políticas públicas).",
            "150 - 250 caracteres",
            "Pendiente"
        ],
        [
            "RES_LINE_01",
            "Línea de Investigación 1 - Biología Molecular",
            "Título: Patogénesis Viral y Oncogénesis por Tax/HBZ\nObjetivo: Comprender los mecanismos moleculares de transformación linfocítica y persistencia proviral en células T humanas.\nLíder: Laboratorio de Virología Molecular IMTAvH / Institut Pasteur",
            "",
            "Nombre de la línea de investigación, resumen de objetivos y centros participantes.",
            "Resumen: 180 - 250 caracteres",
            "Pendiente"
        ],
        [
            "RES_LINE_02",
            "Línea de Investigación 2 - Mielopatía HAM/TSP",
            "Título: Inmunopatogenia y Biomarcadores en HAM/TSP\nObjetivo: Identificación de marcadores inflamatorios en líquido cefalorraquídeo y suero para predicción temprana de deterioro motor.\nLíder: Unidad de Neurología y Neurovirología",
            "",
            "Nombre de la línea de investigación, resumen de objetivos y centros participantes.",
            "Resumen: 180 - 250 caracteres",
            "Pendiente"
        ],
        [
            "RES_LINE_03",
            "Línea de Investigación 3 - Coinfecciones Endémicas",
            "Título: Coinfección HTLV-1, Tuberculosis y Strongyloides\nObjetivo: Evaluación del impacto de la desregulación inmunológica mediada por HTLV-1 en la respuesta inmune contra patógenos oportunistas.\nLíder: IMTAvH - UPCH",
            "",
            "Nombre de la línea de investigación, resumen de objetivos y centros participantes.",
            "Resumen: 180 - 250 caracteres",
            "Pendiente"
        ],
        [
            "RES_LINE_04",
            "Línea de Investigación 4 - Prevención Materno-Infantil",
            "Título: Vigilancia Prenatal y Políticas de Lactancia Segura\nObjetivo: Implementación y modelamiento de costo-efectividad del tamizaje serológico universal en gestantes y provisión de sucedáneos lácteos.\nLíder: Consorcio de Salud Pública",
            "",
            "Nombre de la línea de investigación, resumen de objetivos y centros participantes.",
            "Resumen: 180 - 250 caracteres",
            "Pendiente"
        ],
        [
            "RES_PROJ_01",
            "Proyecto Activo 1",
            "Título: Estudio de Cohorte Longitudinal de Portadores Asintomáticos de HTLV-1 en Lima (2020-2027)\nInvestigador Principal: Dr. / Dra. [Nombre]\nFinanciamiento: [Agencia / Fondo Concursable]\nEstado: En Ejecución\nResumen: Monitoreo semestral de carga proviral y citocinas proinflamatorias en 450 pacientes.",
            "",
            "Título formal del proyecto, PI, entidad financiadora, código del proyecto y breve sinopsis.",
            "Completar todos los campos del proyecto",
            "Pendiente"
        ],
        [
            "RES_PROJ_02",
            "Proyecto Activo 2",
            "Título: Ensayo Clínico Multicéntrico Fase II de Terapia Inmunomoduladora en HAM/TSP Progresiva\nInvestigador Principal: [Nombre]\nFinanciamiento: [Agencia]\nEstado: En Reclutamiento\nResumen: Evaluación de eficacia en la escala de discapacidad motora de Osame a 48 semanas.",
            "",
            "Título formal del proyecto, PI, entidad financiadora, código del proyecto y breve sinopsis.",
            "Completar todos los campos del proyecto",
            "Pendiente"
        ],
        [
            "RES_PUB_01",
            "Publicación Clave 1",
            "Título: Global epidemiology of HTLV-1 and HTLV-2: A systematic review and meta-analysis.\nAutores: Gotuzzo E, Taylor GP, et al.\nRevista: Lancet Infectious Diseases | DOI: 10.1016/S1473-3099(xx)xxxx | Año: 2024",
            "",
            "Cita bibliográfica en formato Vancouver o APA con DOI exacto para indexación automática.",
            "Formato de paper científico estándar",
            "Pendiente"
        ],
        [
            "RES_PUB_02",
            "Publicación Clave 2",
            "Título: Proviral load dynamics and risk of neurological disease in asymptomatic HTLV-1 carriers.\nAutores: Gonzalez E, et al.\nRevista: Journal of Clinical Virology | DOI: 10.1016/j.jcv.2023.xxxx | Año: 2023",
            "",
            "Cita bibliográfica en formato Vancouver o APA con DOI exacto.",
            "Formato de paper científico estándar",
            "Pendiente"
        ]
    ]
    create_content_sheet("04_INVESTIGACION", "LÍNEAS DE INVESTIGACIÓN & PUBLICACIONES", "Líneas de investigación activas, proyectos con financiamiento, ensayos clínicos y artículos indexados.", res_rows)

    # -------------------------------------------------------------
    # HOJA 5: RECURSOS & MULTIMEDIA (RESOURCES)
    # -------------------------------------------------------------
    rec_rows = [
        [
            "REC_HEADER_TITLE",
            "Cabecera - Título de Recursos",
            "Centro Multimedia, Herramientas Interactivas y Guías Clínicas",
            "",
            "Título del hub de recursos educativos y clínicos.",
            "Máx. 70 caracteres",
            "Pendiente"
        ],
        [
            "REC_DOC_01",
            "Documento Descargable 1 - Guía MINSA/UPCH",
            "Título: Guía de Práctica Clínica para el Diagnóstico y Manejo de la Infección por HTLV-1\nPúblico: Médicos Infectólogos, Neurólogos y Médicos Generales\nFormato: PDF (2.4 MB)\nDescripción: Protocolo estandarizado de tamizaje serológico, evaluación de carga proviral y algoritmo terapéutico.",
            "",
            "Título oficial de la guía, autores, formato, peso del PDF y público objetivo.",
            "Descripción: 150 - 220 caracteres",
            "Pendiente"
        ],
        [
            "REC_DOC_02",
            "Documento Descargable 2 - Folleto Pacientes",
            "Título: Manual Informativo para Pacientes y Familias Seropositivas a HTLV-1\nPúblico: Pacientes y Comunidad\nFormato: PDF (1.1 MB)\nDescripción: Lenguaje claro y empático sobre prevención de transmisión, vida saludable y chequeos médicos periódicos.",
            "",
            "Folleto educativo para pacientes seropositivos.",
            "Descripción: 150 - 220 caracteres",
            "Pendiente"
        ],
        [
            "REC_DOC_03",
            "Documento Descargable 3 - Protocolo Laboratorio",
            "Título: Manual de Procedimientos Operativos para Cuantificación de Carga Proviral por qPCR\nPúblico: Biólogos Moleculares y Tecnólogos Médicos\nFormato: PDF (3.8 MB)\nDescripción: Protocolos de extracción de ADN genómico en PBMC y diseño de primers/sondas para gen Tax y albúmina.",
            "",
            "Manual técnico de laboratorio molecular.",
            "Descripción: 150 - 220 caracteres",
            "Pendiente"
        ],
        [
            "REC_VID_01",
            "Video Educativo 1 (YouTube Integrado)",
            "Título: Epidemiología y Manejo Clínico del Virus HTLV-1\nURL / YouTube ID: https://www.youtube.com/watch?v=6Aa3aFccue8 (6Aa3aFccue8)\nCanal / Ponente: Ponencia Científica de Referencia\nCategoría: Epidemiología & Diagnóstico",
            "",
            "Confirmar el enlace oficial de YouTube de la charla o ponencia a incrustar.",
            "URL de YouTube válida",
            "Pendiente"
        ],
        [
            "REC_VID_QUIZ_01",
            "Evaluación Interactiva - Pregunta 1",
            "Pregunta: ¿Cuál es la vía de transmisión del HTLV-1 que presenta mayor riesgo de desarrollar Leucemia/Linfoma de Células T del Adulto (ATL)?\nOpciones:\nA) Transfusión sanguínea\nB) Lactancia materna prolongada (Respuesta Correcta)\nC) Transmisión sexual\nD) Contacto casual de piel\nExplicación: La transmisión vertical temprana mediante lactancia materna prolongada (> 6 meses) es el principal factor de riesgo para el desarrollo de ATL en la edad adulta.",
            "",
            "Pregunta de autoevaluación médica para el video, 4 alternativas, indicar la correcta y retroalimentación.",
            "Formato de opción múltiple estándar",
            "Pendiente"
        ],
        [
            "REC_CALC_INFO",
            "Calculadora de Progresión Clínica",
            "Herramienta interactiva de estratificación de riesgo basada en: Carga Proviral (> 1%), Edad, Sexo, Síntomas neurológicos tempranos, Dermatitis previa y Antecedente familiar.\nPuntaje: 0 a 100 pts con recomendaciones de seguimiento semestral o anual.",
            "",
            "Validar si los criterios y rangos de puntos coinciden con los consensos del IMTAvH.",
            "Revisión por comité clínico",
            "Pendiente"
        ],
        [
            "REC_FORM_INFO",
            "Formulario de Notificación Referencial",
            "Campos del formulario: Rango de edad, Región de procedencia, Tipo serológico (HTLV-1 / HTLV-2), Estado clínico (Asintomático, HAM/TSP, ATL), Disponibilidad de Carga Proviral, Aceptación de términos según Ley N° 29733.",
            "",
            "Verificar si se requieren campos adicionales de epidemiología o contacto.",
            "Formulario de registro voluntario",
            "Pendiente"
        ]
    ]
    create_content_sheet("05_RECURSOS_MULTIMEDIA", "RECURSOS EDUCATIVOS & MULTIMEDIA", "Guías clínicas descargables en PDF, videos de YouTube con preguntas interactivas y herramientas clínicas.", rec_rows)

    # -------------------------------------------------------------
    # HOJA 6: OPORTUNIDADES (OPPORTUNITIES)
    # -------------------------------------------------------------
    opp_rows = [
        [
            "OPP_HEADER_TITLE",
            "Cabecera - Título de Convocatorias",
            "Oportunidades Académicas, Becas y Pasantías de Investigación",
            "",
            "Título principal de la sección de becas y convocatorias.",
            "Máx. 75 caracteres",
            "Pendiente"
        ],
        [
            "OPP_HEADER_DESC",
            "Cabecera - Descripción",
            "Fomentando la formación de la nueva generación de virólogos, infectólogos y epidemiólogos a través de estancias en laboratorios de la red RIII-HTLV.",
            "",
            "Objetivo de capacitación y desarrollo de talento científico joven.",
            "150 - 250 caracteres",
            "Pendiente"
        ],
        [
            "OPP_ITEM_01",
            "Convocatoria 1 - Estancia de Investigación",
            "Título: Pasantía de Entrenamiento en Diagnóstico Molecular de Retrovirología (UPCH - Lima)\nDirigido a: Estudiantes de posgrado en Ciencias Biomédicas y Residentes de Infectología\nDuración: 3 a 6 meses\nRequisitos: Grado de bachiller o médico cirujano, experiencia básica en biología molecular y carta de motivación.\nFecha Límite: [Fecha oficial]\nContacto: [Correo institucional]",
            "",
            "Detalles completos de la pasantía o rotación clínica/laboratorial.",
            "Completar todos los campos del formato",
            "Pendiente"
        ],
        [
            "OPP_ITEM_02",
            "Convocatoria 2 - Beca Posdoctoral",
            "Título: Beca de Investigación Posdoctoral en Inmunología de Linfocitos T y HTLV-1\nInstitución Anfitriona: Imperial College London / IMTAvH\nFinanciamiento: Estipendio mensual completo + fondos para reactivos\nFecha Límite: [Fecha oficial]\nContacto: [Correo institucional]",
            "",
            "Detalles de beca o financiamiento para proyectos posdoctorales.",
            "Completar todos los campos del formato",
            "Pendiente"
        ]
    ]
    create_content_sheet("06_OPORTUNIDADES", "OPORTUNIDADES & BECAS", "Convocatorias abiertas para pasantías, becas de posgrado y estancias internacionales en laboratorios de la red.", opp_rows)

    # -------------------------------------------------------------
    # HOJA 7: CONTACTO & LEGAL (CONTACT)
    # -------------------------------------------------------------
    contact_rows = [
        [
            "CONT_HEADER_TITLE",
            "Cabecera - Título de Contacto",
            "Canales Oficiales de Comunicación y Sede Central",
            "",
            "Título de la sección de contacto institucional.",
            "Máx. 65 caracteres",
            "Pendiente"
        ],
        [
            "CONT_MAIN_SEAT",
            "Sede Central de Coordinación",
            "Institución: Instituto de Medicina Tropical Alexander von Humboldt (IMTAvH)\nUniversidad: Universidad Peruana Cayetano Heredia (UPCH)\nDirección: Av. Honorio Delgado 430, San Martín de Porres, Lima, Perú\nCódigo Postal: 15102",
            "",
            "Dirección física oficial y nombre del pabellón o laboratorio sede.",
            "Dirección postal completa",
            "Pendiente"
        ],
        [
            "CONT_EMAIL_OFFICIAL",
            "Correos Electrónicos Oficiales",
            "Coordinación General: riii.htlv@upch.pe\nConsultas Científicas: imtavh.investigacion@oficinas-upch.pe\nSoporte Web / Registro: contacto@htlv-network.org",
            "",
            "Listado de correos electrónicos institucionales oficiales para contacto público.",
            "Correos válidos con dominio oficial",
            "Pendiente"
        ],
        [
            "CONT_PHONE",
            "Teléfonos Institucionales",
            "Central Telefónica: +51 (1) 319-0000 anexo 201300\nAtención IMTAvH: +51 (1) 482-3903",
            "",
            "Números telefónicos oficiales con código de país (+51) y anexos.",
            "Teléfono con código internacional",
            "Pendiente"
        ],
        [
            "CONT_SOCIAL_MEDIA",
            "Redes Sociales y Canales Oficiales",
            "LinkedIn: [URL oficial]\nX / Twitter: [URL oficial]\nYouTube: [URL oficial]\nResearchGate / ORCID: [URL oficial]",
            "",
            "Enlaces a las redes académicas e institucionales oficiales de la red.",
            "URLs completas",
            "Pendiente"
        ],
        [
            "LEGAL_PRIVACY_POLICY",
            "Aviso de Privacidad y Consentimiento (Ley N° 29733)",
            "Texto legal de tratamiento de datos personales conforme a la Ley N° 29733 (Ley de Protección de Datos Personales de la República del Perú) para formularios, registros de la red y envío de boletines informativos.",
            "",
            "Aprobado por el área legal o asesoría de bioética institucional.",
            "Texto legal formal",
            "Pendiente"
        ],
        [
            "LEGAL_DISCLAIMER",
            "Descargo de Responsabilidad Médica",
            "La información contenida en esta plataforma es de carácter exclusivamente científico, educativo e informativo. No constituye asesoramiento médico personalizado ni reemplaza el diagnóstico o tratamiento presencial por un profesional médico especialista colegiado.",
            "",
            "Texto legal obligatorio para sitios web con contenido biomédico.",
            "180 - 300 caracteres",
            "Pendiente"
        ]
    ]
    create_content_sheet("07_CONTACTO_Y_LEGAL", "CONTACTO & TÉRMINOS LEGALES", "Dirección de la sede central, correos oficiales, redes sociales, aviso de privacidad (Ley 29733) y descargo médico.", contact_rows)

    # Save workbook
    wb.save(output_path)
    print(f"Workbook successfully created at: {output_path}")

if __name__ == "__main__":
    import os
    target = os.path.abspath("RECOLECCION_CONTENIDO_OFICIAL_WEB_HTLV.xlsx")
    build_content_collection_workbook(target)
